import base64
import csv
import hashlib
import io
import json
import re
import uuid
from collections import defaultdict
from dataclasses import dataclass
from datetime import UTC, datetime
from decimal import Decimal
from xml.sax.saxutils import escape

from sqlalchemy.orm import Session

from app.core.settings import Settings
from app.models.debate_message import DebateMessage
from app.models.document_chunk import DocumentChunk
from app.models.evaluation_result import EvaluationResult
from app.models.evidence_item import EvidenceItem
from app.models.hypothesis_candidate import HypothesisCandidate
from app.models.hypothesis_evidence_link import HypothesisEvidenceLink
from app.models.hypothesis_version import HypothesisVersion
from app.models.judge_verdict import JudgeVerdict
from app.models.research_feedback import ResearchFeedback
from app.models.research_input_item import ResearchInputItem
from app.models.research_run import ResearchRun
from app.models.research_run_event import ResearchRunEvent
from app.models.session_file import SessionFile
from app.models.user import User
from app.models.user_agent import UserAgent
from app.models.enums import (
    DebateRole,
    DocumentProcessingStatus,
    ResearchFeedbackTarget,
    ResearchInputKind,
    ResearchRunStatus,
    ResearchRunStage,
    ResearchRunTrigger,
)
from app.repositories.research_repository import ResearchRepository
from app.schemas.research import (
    DebateMessageResponse,
    EvaluationResultResponse,
    EvidenceItemResponse,
    HypothesisCandidateResponse,
    HypothesisEvidenceLinkResponse,
    HypothesisVersionResponse,
    JudgeVerdictResponse,
    PipelineSettingsResponse,
    ResearchExportResponse,
    ResearchFeedbackCreateRequest,
    ResearchFeedbackListResponse,
    ResearchFeedbackResponse,
    ResearchGraphEdgeResponse,
    ResearchGraphNodeResponse,
    ResearchGraphResponse,
    ResearchInputRequest,
    ResearchInputResponse,
    ResearchRunDetailResponse,
    ResearchRunEventResponse,
    ResearchRunListResponse,
    ResearchRunProgressResponse,
    ResearchRunSummaryResponse,
)
from app.services.exceptions import ConflictError, NotFoundError, ServiceError, ValidationError
from app.services.document_text_extractor import DocumentTextExtractor
from app.services.research_llm_orchestrator import EvidenceDraft, HypothesisDraft, ResearchLlmOrchestrator
from app.services.research_retrieval_service import ResearchRetrievalService, RetrievalResult

_CHUNK_MAX_CHARS = 1800
_WORD_RE = re.compile(r"\s+")


@dataclass(frozen=True)
class NormalizedInput:
    kind: ResearchInputKind
    label: str
    text: str
    position: int


class ResearchService:
    def __init__(
        self,
        *,
        session: Session,
        repository: ResearchRepository,
        settings: Settings,
        text_extractor: DocumentTextExtractor | None = None,
        retrieval_service: ResearchRetrievalService | None = None,
        llm_orchestrator: ResearchLlmOrchestrator | None = None,
    ) -> None:
        self._session = session
        self._repository = repository
        self._settings = settings
        self._text_extractor = text_extractor or DocumentTextExtractor()
        self._retrieval_service = retrieval_service
        self._llm_orchestrator = llm_orchestrator

    def list_runs(self, *, user: User, chat_session_id: uuid.UUID) -> ResearchRunListResponse:
        chat_session = self._get_chat_session(user=user, chat_session_id=chat_session_id)
        runs = self._repository.list_runs(self._session, user_id=user.id, chat_session_id=chat_session_id)
        return ResearchRunListResponse(
            items=[ResearchRunSummaryResponse.model_validate(run) for run in runs],
            active_run_id=chat_session.active_research_run_id,
        )

    def get_pipeline_settings(self) -> PipelineSettingsResponse:
        settings = self._repository.get_pipeline_settings(self._session)
        return PipelineSettingsResponse(
            default_hypothesis_count=settings.default_hypothesis_count,
            debate_round_limit=settings.debate_round_limit,
            retrieval_limit=settings.retrieval_limit,
            evaluator_weights=settings.evaluator_weights,
            excluded_directions=settings.excluded_directions,
            domain_constraints=settings.domain_constraints,
        )

    def get_active_run(self, *, user: User, chat_session_id: uuid.UUID) -> ResearchRunDetailResponse:
        self._get_chat_session(user=user, chat_session_id=chat_session_id)
        run = self._repository.get_active_run(self._session, user_id=user.id, chat_session_id=chat_session_id)
        if run is None:
            raise NotFoundError("Активная версия исследования не найдена.")
        return self._build_detail_response(run)

    def get_run(self, *, user: User, chat_session_id: uuid.UUID, run_id: uuid.UUID) -> ResearchRunDetailResponse:
        run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        return self._build_detail_response(run)

    def get_progress(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
    ) -> ResearchRunProgressResponse:
        run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        events = self._repository.list_run_events(self._session, run_id=run.id)
        current_stage = events[-1].stage if events else None
        progress_percent = events[-1].progress_percent if events else 0
        return ResearchRunProgressResponse(
            run=ResearchRunSummaryResponse.model_validate(run),
            current_stage=current_stage,
            progress_percent=progress_percent,
            events=[ResearchRunEventResponse.model_validate(event) for event in events],
        )

    def activate_run(self, *, user: User, chat_session_id: uuid.UUID, run_id: uuid.UUID) -> ResearchRunSummaryResponse:
        try:
            chat_session = self._get_chat_session(user=user, chat_session_id=chat_session_id, for_update=True)
            run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
            if run.status != ResearchRunStatus.COMPLETED:
                raise ConflictError("Можно переключаться только на завершённые версии исследования.")
            chat_session.active_research_run_id = run.id
            self._session.commit()
            return ResearchRunSummaryResponse.model_validate(run)
        except Exception:
            self._session.rollback()
            raise

    def create_run(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        input_requests: list[ResearchInputRequest],
        hypothesis_count: int,
    ) -> ResearchRunDetailResponse:
        run = self.create_run_record(
            user=user,
            chat_session_id=chat_session_id,
            input_requests=input_requests,
            hypothesis_count=hypothesis_count,
        )
        return self.execute_run_pipeline(user=user, chat_session_id=chat_session_id, run_id=run.id)

    def create_run_record(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        input_requests: list[ResearchInputRequest],
        hypothesis_count: int,
    ) -> ResearchRunDetailResponse:
        normalized_inputs = self._normalize_inputs(input_requests)
        run = self._create_run_record_from_inputs(
            user=user,
            chat_session_id=chat_session_id,
            normalized_inputs=normalized_inputs,
            hypothesis_count=hypothesis_count,
            trigger=ResearchRunTrigger.INITIAL,
            parent_run_id=None,
        )
        return self._build_detail_response(run)

    def regenerate_run(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        hypothesis_count: int | None,
    ) -> ResearchRunDetailResponse:
        source_run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        if source_run.status == ResearchRunStatus.RUNNING:
            raise ConflictError("Нельзя перегенерировать версию, пока она выполняется.")
        source_inputs = self._repository.list_input_items(self._session, run_id=source_run.id)
        normalized_inputs = [
            NormalizedInput(kind=item.kind, label=item.label, text=item.text, position=item.position)
            for item in source_inputs
        ]
        run = self._create_run_record_from_inputs(
            user=user,
            chat_session_id=chat_session_id,
            normalized_inputs=normalized_inputs,
            hypothesis_count=hypothesis_count or source_run.hypothesis_count,
            trigger=ResearchRunTrigger.REGENERATE,
            parent_run_id=source_run.id,
        )
        return self.execute_run_pipeline(user=user, chat_session_id=chat_session_id, run_id=run.id)

    def regenerate_run_record(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        hypothesis_count: int | None,
    ) -> ResearchRunDetailResponse:
        source_run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        if source_run.status == ResearchRunStatus.RUNNING:
            raise ConflictError("Нельзя перегенерировать версию, пока она выполняется.")
        source_inputs = self._repository.list_input_items(self._session, run_id=source_run.id)
        normalized_inputs = [
            NormalizedInput(kind=item.kind, label=item.label, text=item.text, position=item.position)
            for item in source_inputs
        ]
        run = self._create_run_record_from_inputs(
            user=user,
            chat_session_id=chat_session_id,
            normalized_inputs=normalized_inputs,
            hypothesis_count=hypothesis_count or source_run.hypothesis_count,
            trigger=ResearchRunTrigger.REGENERATE,
            parent_run_id=source_run.id,
        )
        return self._build_detail_response(run)

    def edit_run(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        input_requests: list[ResearchInputRequest],
        hypothesis_count: int | None,
    ) -> ResearchRunDetailResponse:
        source_run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        if source_run.status == ResearchRunStatus.RUNNING:
            raise ConflictError("Нельзя редактировать версию, пока она выполняется.")
        run = self._create_run_record_from_inputs(
            user=user,
            chat_session_id=chat_session_id,
            normalized_inputs=self._normalize_inputs(input_requests),
            hypothesis_count=hypothesis_count or source_run.hypothesis_count,
            trigger=ResearchRunTrigger.EDIT,
            parent_run_id=source_run.id,
        )
        return self.execute_run_pipeline(user=user, chat_session_id=chat_session_id, run_id=run.id)

    def edit_run_record(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        input_requests: list[ResearchInputRequest],
        hypothesis_count: int | None,
    ) -> ResearchRunDetailResponse:
        source_run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        if source_run.status == ResearchRunStatus.RUNNING:
            raise ConflictError("Нельзя редактировать версию, пока она выполняется.")
        run = self._create_run_record_from_inputs(
            user=user,
            chat_session_id=chat_session_id,
            normalized_inputs=self._normalize_inputs(input_requests),
            hypothesis_count=hypothesis_count or source_run.hypothesis_count,
            trigger=ResearchRunTrigger.EDIT,
            parent_run_id=source_run.id,
        )
        return self._build_detail_response(run)

    def cancel_run(self, *, user: User, chat_session_id: uuid.UUID, run_id: uuid.UUID) -> ResearchRunSummaryResponse:
        try:
            run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
            if run.status != ResearchRunStatus.RUNNING:
                raise ConflictError("Можно отменить только выполняющуюся версию исследования.")
            run.status = ResearchRunStatus.CANCELLED
            run.completed_at = datetime.now(UTC)
            self._append_event(
                run=run,
                stage=ResearchRunStage.CANCELLED,
                progress_percent=100,
                message="Исследовательский запуск отменён.",
            )
            self._session.commit()
            return ResearchRunSummaryResponse.model_validate(run)
        except Exception:
            self._session.rollback()
            raise

    def create_feedback(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        payload: ResearchFeedbackCreateRequest,
    ) -> ResearchFeedbackResponse:
        try:
            run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
            self._validate_feedback_target(run=run, payload=payload)
            feedback = self._repository.create_feedback(
                self._session,
                ResearchFeedback(
                    run_id=run.id,
                    user_id=user.id,
                    target_type=payload.target_type,
                    target_id=payload.target_id,
                    outcome=payload.outcome,
                    rating=payload.rating,
                    comment=self._normalize_optional_text(payload.comment),
                    correction=self._normalize_optional_text(payload.correction),
                    feedback_metadata=payload.metadata,
                ),
            )
            self._session.commit()
            return ResearchFeedbackResponse.model_validate(feedback)
        except Exception:
            self._session.rollback()
            raise

    def list_feedback(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
    ) -> ResearchFeedbackListResponse:
        run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        return ResearchFeedbackListResponse(
            items=[
                ResearchFeedbackResponse.model_validate(item)
                for item in self._repository.list_feedback(self._session, run_id=run.id)
            ]
        )

    def export_run(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        export_format: str,
    ) -> ResearchExportResponse:
        detail = self.get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        normalized_format = export_format.strip().lower()
        if normalized_format == "json":
            return ResearchExportResponse(
                filename=f"research-run-{detail.version_number}.json",
                content_type="application/json",
                content=detail.model_dump_json(indent=2),
            )
        if normalized_format in {"md", "markdown"}:
            return ResearchExportResponse(
                filename=f"research-run-{detail.version_number}.md",
                content_type="text/markdown; charset=utf-8",
                content=self._build_markdown_export(detail),
            )
        if normalized_format == "csv":
            return ResearchExportResponse(
                filename=f"research-run-{detail.version_number}.csv",
                content_type="text/csv; charset=utf-8",
                content=self._build_csv_export(detail),
            )
        if normalized_format == "docx":
            return ResearchExportResponse(
                filename=f"research-run-{detail.version_number}.docx",
                content_type="application/vnd.openxmlformats-officedocument.wordprocessingml.document",
                content=self._build_docx_export(detail),
                content_encoding="base64",
            )
        if normalized_format == "pdf":
            return ResearchExportResponse(
                filename=f"research-run-{detail.version_number}.pdf",
                content_type="application/pdf",
                content=self._build_pdf_export(detail),
                content_encoding="base64",
            )
        raise ValidationError("Неподдерживаемый формат экспорта.")

    def get_graph(self, *, user: User, chat_session_id: uuid.UUID, run_id: uuid.UUID) -> ResearchGraphResponse:
        run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
        detail = self._build_detail_response(run)
        nodes: list[ResearchGraphNodeResponse] = [
            ResearchGraphNodeResponse(
                id=f"run:{run.id}",
                type="run",
                label=f"Версия {run.version_number}",
                description=run.title,
                metadata={"status": run.status.value},
            ),
        ]
        edges: list[ResearchGraphEdgeResponse] = []

        for input_item in detail.inputs:
            input_node_id = f"input:{input_item.id}"
            nodes.append(
                ResearchGraphNodeResponse(
                    id=input_node_id,
                    type="input",
                    label=input_item.label,
                    description=input_item.text,
                    metadata={"kind": input_item.kind.value},
                )
            )
            edges.append(
                ResearchGraphEdgeResponse(
                    id=f"{input_node_id}->run:{run.id}",
                    source=input_node_id,
                    target=f"run:{run.id}",
                    type="defines",
                    label="задаёт рамку",
                )
            )

        source_files = self._repository.list_active_files(
            self._session,
            user_id=run.user_id,
            chat_session_id=run.chat_session_id,
        )
        for source_file in source_files:
            file_node_id = f"file:{source_file.id}"
            nodes.append(
                ResearchGraphNodeResponse(
                    id=file_node_id,
                    type="file",
                    label=source_file.original_filename,
                    description=source_file.processing_status.value,
                    metadata={
                        "session_file_id": str(source_file.id),
                        "processing_status": source_file.processing_status.value,
                        "processing_error": source_file.processing_error,
                    },
                )
            )
            edges.append(
                ResearchGraphEdgeResponse(
                    id=f"run:{run.id}->file:{source_file.id}",
                    source=f"run:{run.id}",
                    target=file_node_id,
                    type="uses_source",
                    label="использует источник",
                )
            )
            for chunk in self._repository.list_chunks_for_file(self._session, session_file_id=source_file.id):
                chunk_node_id = f"chunk:{chunk.id}"
                if all(node.id != chunk_node_id for node in nodes):
                    nodes.append(
                        ResearchGraphNodeResponse(
                            id=chunk_node_id,
                            type="chunk",
                            label=f"Фрагмент {chunk.position + 1}",
                            description=chunk.content,
                            metadata={
                                "session_file_id": str(source_file.id),
                                "chunk_id": str(chunk.id),
                                "chunk_position": chunk.position,
                            },
                        )
                    )
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"file:{source_file.id}->chunk:{chunk.id}",
                        source=file_node_id,
                        target=chunk_node_id,
                        type="has_chunk",
                        label="разбит на фрагмент",
                    )
                )

        for evidence in detail.evidence:
            chunk_node_id = f"chunk:{evidence.chunk_id}" if evidence.chunk_id else None
            if chunk_node_id is not None and all(node.id != chunk_node_id for node in nodes):
                nodes.append(
                    ResearchGraphNodeResponse(
                        id=chunk_node_id,
                        type="chunk",
                        label="Фрагмент источника",
                        description=evidence.quote or evidence.summary,
                        metadata={
                            "session_file_id": str(evidence.session_file_id) if evidence.session_file_id else None,
                            "chunk_id": str(evidence.chunk_id),
                        },
                    )
                )
            evidence_node_id = f"evidence:{evidence.id}"
            nodes.append(
                ResearchGraphNodeResponse(
                    id=evidence_node_id,
                    type="evidence",
                    label=evidence.title,
                    description=evidence.summary,
                    metadata={
                        "kind": evidence.kind.value,
                        "confidence": float(evidence.confidence),
                        "session_file_id": str(evidence.session_file_id) if evidence.session_file_id else None,
                        "chunk_id": str(evidence.chunk_id) if evidence.chunk_id else None,
                    },
                )
            )
            edges.append(
                ResearchGraphEdgeResponse(
                    id=f"run:{run.id}->evidence:{evidence.id}",
                    source=f"run:{run.id}",
                    target=evidence_node_id,
                    type="contains",
                )
            )
            if chunk_node_id is not None:
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"{chunk_node_id}->evidence:{evidence.id}",
                        source=chunk_node_id,
                        target=evidence_node_id,
                        type="supports_trace",
                        label="источник evidence",
                    )
                )

        for hypothesis in detail.hypotheses:
            hypothesis_node_id = f"hypothesis:{hypothesis.id}"
            nodes.append(
                ResearchGraphNodeResponse(
                    id=hypothesis_node_id,
                    type="hypothesis",
                    label=hypothesis.title,
                    description=hypothesis.statement,
                    metadata={"position": hypothesis.position},
                )
            )
            edges.append(
                ResearchGraphEdgeResponse(
                    id=f"run:{run.id}->hypothesis:{hypothesis.id}",
                    source=f"run:{run.id}",
                    target=hypothesis_node_id,
                    type="generates",
                )
            )
            for link in hypothesis.evidence_links:
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"evidence:{link.evidence_id}->hypothesis:{hypothesis.id}:{link.relation.value}",
                        source=f"evidence:{link.evidence_id}",
                        target=hypothesis_node_id,
                        type=link.relation.value,
                        label=link.rationale,
                    )
                )
            for version in hypothesis.versions:
                version_node_id = f"hypothesis-version:{version.id}"
                nodes.append(
                    ResearchGraphNodeResponse(
                        id=version_node_id,
                        type="hypothesis_version",
                        label=f"Версия {version.version_number}",
                        description=version.statement,
                        metadata={
                            "hypothesis_id": str(hypothesis.id),
                            "created_by_role": version.created_by_role.value if version.created_by_role else None,
                        },
                    )
                )
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"hypothesis:{hypothesis.id}->hypothesis-version:{version.id}",
                        source=hypothesis_node_id,
                        target=version_node_id,
                        type="has_version",
                        label=version.change_summary,
                    )
                )
            for message in hypothesis.debate_messages:
                message_node_id = f"debate-message:{message.id}"
                nodes.append(
                    ResearchGraphNodeResponse(
                        id=message_node_id,
                        type="debate_message",
                        label=f"{message.role.value} / round {message.round_number}",
                        description=message.content,
                        metadata={"hypothesis_id": str(hypothesis.id)},
                    )
                )
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"debate-message:{message.id}->hypothesis:{hypothesis.id}",
                        source=message_node_id,
                        target=hypothesis_node_id,
                        type="critiques",
                    )
                )
            for evaluation in hypothesis.evaluations:
                evaluation_node_id = f"evaluation:{evaluation.id}"
                nodes.append(
                    ResearchGraphNodeResponse(
                        id=evaluation_node_id,
                        type="evaluation",
                        label=evaluation.evaluator_name,
                        description=evaluation.verdict,
                        metadata={"score": float(evaluation.score), "hypothesis_id": str(hypothesis.id)},
                    )
                )
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"hypothesis:{hypothesis.id}->evaluation:{evaluation.id}",
                        source=hypothesis_node_id,
                        target=evaluation_node_id,
                        type="evaluated_by",
                    )
                )

        if detail.verdict is not None:
            verdict_node_id = f"verdict:{detail.verdict.id}"
            nodes.append(
                ResearchGraphNodeResponse(
                    id=verdict_node_id,
                    type="verdict",
                    label="Вердикт судьи",
                    description=detail.verdict.recommendation,
                    metadata={"ranking": detail.verdict.ranking},
                )
            )
            edges.append(
                ResearchGraphEdgeResponse(
                    id=f"run:{run.id}->verdict:{detail.verdict.id}",
                    source=f"run:{run.id}",
                    target=verdict_node_id,
                    type="synthesizes",
                )
            )
            for index, check in enumerate(detail.verdict.next_checks, start=1):
                check_node_id = f"next-check:{detail.verdict.id}:{index}"
                nodes.append(
                    ResearchGraphNodeResponse(
                        id=check_node_id,
                        type="next_check",
                        label=f"Проверка {index}",
                        description=check,
                    )
                )
                edges.append(
                    ResearchGraphEdgeResponse(
                        id=f"verdict:{detail.verdict.id}->next-check:{index}",
                        source=verdict_node_id,
                        target=check_node_id,
                        type="recommends",
                    )
                )

        return ResearchGraphResponse(run_id=run.id, nodes=nodes, edges=edges)

    def _create_run_record_from_inputs(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        normalized_inputs: list[NormalizedInput],
        hypothesis_count: int,
        trigger: ResearchRunTrigger,
        parent_run_id: uuid.UUID | None,
    ) -> ResearchRun:
        try:
            chat_session = self._get_chat_session(user=user, chat_session_id=chat_session_id, for_update=True)
            if self._repository.has_running_run(self._session, chat_session_id=chat_session_id):
                raise ConflictError("В этом чате уже выполняется исследовательский запуск.")
            pipeline_settings = self._repository.get_pipeline_settings(self._session)
            effective_hypothesis_count = hypothesis_count or pipeline_settings.default_hypothesis_count

            now = datetime.now(UTC)
            run = self._repository.create_run(
                self._session,
                ResearchRun(
                    chat_session_id=chat_session.id,
                    user_id=user.id,
                    parent_run_id=parent_run_id,
                    version_number=self._repository.next_version_number(self._session, chat_session_id=chat_session.id),
                    trigger=trigger,
                    status=ResearchRunStatus.RUNNING,
                    title=self._create_title(normalized_inputs),
                    input_hash=self._create_input_hash(normalized_inputs),
                    model_name=self._settings.model_provider_model,
                    hypothesis_count=effective_hypothesis_count,
                    started_at=now,
                ),
            )
            for item in normalized_inputs:
                self._repository.create_input_item(
                    self._session,
                    ResearchInputItem(
                        run_id=run.id,
                        kind=item.kind,
                        label=item.label,
                        text=item.text,
                        position=item.position,
                    ),
                )
            self._append_event(
                run=run,
                stage=ResearchRunStage.QUEUED,
                progress_percent=0,
                message="Исследовательский запуск создан.",
            )
            self._session.commit()
            return run
        except Exception:
            self._session.rollback()
            raise

    def execute_run_pipeline(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
    ) -> ResearchRunDetailResponse:
        run: ResearchRun | None = None
        try:
            chat_session = self._get_chat_session(user=user, chat_session_id=chat_session_id, for_update=True)
            run = self._get_run(user=user, chat_session_id=chat_session_id, run_id=run_id)
            input_items = self._repository.list_input_items(self._session, run_id=run.id)
            normalized_inputs = [
                NormalizedInput(kind=item.kind, label=item.label, text=item.text, position=item.position)
                for item in input_items
            ]
            pipeline_settings = self._repository.get_pipeline_settings(self._session)
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.INGESTION, progress_percent=10, message="Обработка источников началась.")
            files = self._repository.list_active_files(self._session, user_id=user.id, chat_session_id=chat_session.id)
            chunks = self._prepare_document_chunks(run=run, files=files)
            research_brief = self._research_brief(normalized_inputs)
            research_brief = self._apply_pipeline_settings_to_brief(research_brief, pipeline_settings)
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.RETRIEVAL, progress_percent=25, message="Поиск релевантных фрагментов начался.")
            evidence_drafts, evidence_source_chunks = self._extract_evidence_drafts(
                run=run,
                research_brief=research_brief,
                chunks=chunks,
                files=files,
                retrieval_limit=pipeline_settings.retrieval_limit,
            )
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.EVIDENCE, progress_percent=40, message="Evidence pack сформирован.")
            evidence = self._persist_evidence(
                run=run,
                evidence_drafts=evidence_drafts,
                source_chunks=evidence_source_chunks,
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.EVIDENCE,
                progress_percent=48,
                message=f"Evidence сохранены: {len(evidence)} шт.",
                metadata={"evidence_count": len(evidence)},
            )
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.HYPOTHESIS_GENERATION, progress_percent=55, message="Генерация стартовых гипотез началась.")
            hypothesis_drafts = self._generate_hypothesis_drafts(
                research_brief=research_brief,
                evidence_drafts=evidence_drafts,
                hypothesis_count=run.hypothesis_count,
            )
            hypotheses = self._persist_hypotheses(
                run=run,
                hypothesis_drafts=hypothesis_drafts,
                evidence=evidence,
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.HYPOTHESIS_GENERATION,
                progress_percent=62,
                message=f"Стартовые гипотезы сохранены: {len(hypotheses)} шт.",
                metadata={"hypothesis_count": len(hypotheses)},
            )
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.DEBATE, progress_percent=70, message="Debate loop начался.")
            refined_statements = self._create_debate_artifacts(
                run=run,
                hypothesis_drafts=hypothesis_drafts,
                hypotheses=hypotheses,
                evidence_drafts=evidence_drafts,
                round_limit=pipeline_settings.debate_round_limit,
            )
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.EVALUATION, progress_percent=82, message="Независимая оценка гипотез началась.")
            evaluation_payloads = self._create_evaluation_artifacts(
                run=run,
                chat_session_id=chat_session.id,
                user=user,
                hypothesis_drafts=hypothesis_drafts,
                hypotheses=hypotheses,
                refined_statements=refined_statements,
                evidence_drafts=evidence_drafts,
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.EVALUATION,
                progress_percent=90,
                message=f"Оценки агентов сохранены: {len(evaluation_payloads)} шт.",
                metadata={"evaluation_count": len(evaluation_payloads)},
            )
            self._check_cancelled(run)
            self._append_event_and_commit(run=run, stage=ResearchRunStage.JUDGE, progress_percent=93, message="Финальный судья формирует вердикт.")
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=94,
                message="Судья сравнивает гипотезы между собой.",
                metadata={"judge_step": "compare_hypotheses"},
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=95,
                message="Судья проверяет доказательную базу и связи с evidence.",
                metadata={"judge_step": "check_evidence"},
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=96,
                message="Судья учитывает критику Защитника, Атакующего и Производственника.",
                metadata={"judge_step": "review_debate"},
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=97,
                message="Судья учитывает оценки дополнительных агентов.",
                metadata={"judge_step": "review_evaluations"},
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=98,
                message="Судья готовит итоговый приоритет и первые проверки.",
                metadata={"judge_step": "prepare_ranking"},
            )
            verdict = self._create_judge_verdict(
                run=run,
                research_brief=research_brief,
                hypothesis_drafts=hypothesis_drafts,
                refined_statements=refined_statements,
                evidence_drafts=evidence_drafts,
                evaluation_payloads=evaluation_payloads,
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.JUDGE,
                progress_percent=99,
                message="Вердикт судьи сохранён.",
                metadata={"judge_verdict_id": str(verdict.id)},
            )

            run.status = ResearchRunStatus.COMPLETED
            run.completed_at = datetime.now(UTC)
            chat_session.is_started = True
            if run.trigger != ResearchRunTrigger.REGENERATE:
                chat_session.title = run.title
            chat_session.active_research_run_id = run.id
            self._append_event(run=run, stage=ResearchRunStage.COMPLETED, progress_percent=100, message="Исследовательский запуск завершён.")
            self._session.commit()
            return self._build_detail_response(run)
        except ConflictError as exc:
            self._session.rollback()
            if run is not None:
                self._mark_run_cancelled(user=user, chat_session_id=chat_session_id, run_id=run.id, message=exc.detail)
            raise
        except Exception as exc:
            self._session.rollback()
            if run is not None:
                self._mark_run_failed(user=user, chat_session_id=chat_session_id, run_id=run.id, error=exc)
            raise

    def _prepare_document_chunks(self, *, run: ResearchRun, files: list[SessionFile]) -> list[DocumentChunk]:
        chunks: list[DocumentChunk] = []
        for session_file in files:
            try:
                existing_chunks = self._repository.list_chunks_for_file(self._session, session_file_id=session_file.id)
                if existing_chunks:
                    chunks.extend(existing_chunks)
                    if session_file.processing_status != DocumentProcessingStatus.PROCESSED:
                        self._set_file_processing_status(session_file, DocumentProcessingStatus.PROCESSED)
                        self._append_event_and_commit(
                            run=run,
                            stage=ResearchRunStage.INGESTION,
                            progress_percent=18,
                            message=f"Файл уже подготовлен: {session_file.original_filename}.",
                            metadata={"file_id": str(session_file.id), "filename": session_file.original_filename},
                        )
                    continue

                self._set_file_processing_status(session_file, DocumentProcessingStatus.PARSING)
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.INGESTION,
                    progress_percent=12,
                    message=f"Парсинг файла: {session_file.original_filename}.",
                    metadata={"file_id": str(session_file.id), "filename": session_file.original_filename},
                )
                extracted = self._text_extractor.extract(
                    path=self._settings.uploads_dir / session_file.object_key,
                    kind=session_file.kind,
                    content_type=session_file.content_type,
                )
                if extracted is None or not extracted.text.strip():
                    self._set_file_processing_status(
                        session_file,
                        DocumentProcessingStatus.UNSUPPORTED,
                        error="Для этого формата пока нет текстового парсера.",
                    )
                    self._append_event_and_commit(
                        run=run,
                        stage=ResearchRunStage.INGESTION,
                        progress_percent=16,
                        message=f"Формат файла пока не поддержан: {session_file.original_filename}.",
                        metadata={"file_id": str(session_file.id), "filename": session_file.original_filename},
                    )
                    continue

                for position, content in enumerate(self._split_text(extracted.text)):
                    chunk = self._repository.create_chunk(
                        self._session,
                        DocumentChunk(
                            session_file_id=session_file.id,
                            chat_session_id=session_file.chat_session_id,
                            user_id=session_file.user_id,
                            position=position,
                            content=content,
                            content_hash=hashlib.sha256(content.encode("utf-8")).hexdigest(),
                            token_count=len(content.split()),
                            source_metadata=extracted.metadata | {"original_filename": session_file.original_filename},
                        ),
                    )
                    chunks.append(chunk)
                self._set_file_processing_status(
                    session_file,
                    DocumentProcessingStatus.CHUNKED,
                    text_extracted_at=datetime.now(UTC),
                )
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.INGESTION,
                    progress_percent=20,
                    message=f"Файл разбит на фрагменты: {session_file.original_filename}.",
                    metadata={
                        "file_id": str(session_file.id),
                        "filename": session_file.original_filename,
                        "chunk_count": len([chunk for chunk in chunks if chunk.session_file_id == session_file.id]),
                    },
                )
            except (OSError, UnicodeDecodeError, ValidationError) as exc:
                self._set_file_processing_status(
                    session_file,
                    DocumentProcessingStatus.FAILED,
                    error=str(exc)[:500],
                )
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.INGESTION,
                    progress_percent=16,
                    message=f"Ошибка обработки файла: {session_file.original_filename}.",
                    metadata={"file_id": str(session_file.id), "filename": session_file.original_filename},
                )
        return chunks

    @staticmethod
    def _split_text(text: str) -> list[str]:
        normalized = _WORD_RE.sub(" ", text).strip()
        if not normalized:
            return []
        chunks: list[str] = []
        cursor = 0
        while cursor < len(normalized):
            next_cursor = min(cursor + _CHUNK_MAX_CHARS, len(normalized))
            if next_cursor < len(normalized):
                boundary = normalized.rfind(" ", cursor, next_cursor)
                if boundary > cursor + 400:
                    next_cursor = boundary
            chunks.append(normalized[cursor:next_cursor].strip())
            cursor = next_cursor
        return [chunk for chunk in chunks if chunk]

    def _extract_evidence_drafts(
        self,
        *,
        run: ResearchRun,
        research_brief: str,
        chunks: list[DocumentChunk],
        retrieval_limit: int,
        files: list[SessionFile] | None = None,
    ) -> tuple[list[EvidenceDraft], list[DocumentChunk]]:
        orchestrator = self._require_llm_orchestrator()
        if not chunks:
            return orchestrator.extract_evidence(research_brief=research_brief, fragments=[], limit=8), []
        retrieval_service = self._require_retrieval_service()
        files_with_chunks = self._files_for_chunks(files=files or [], chunks=chunks)
        self._set_files_processing_status(files_with_chunks, DocumentProcessingStatus.INDEXED)
        self._append_event_and_commit(
            run=run,
            stage=ResearchRunStage.RETRIEVAL,
            progress_percent=28,
            message="Фрагменты источников переданы на индексацию.",
            metadata={"file_count": len(files_with_chunks), "chunk_count": len(chunks)},
        )
        try:
            retrieval_service.ensure_chunk_embeddings(chunks)
        except Exception as exc:
            self._set_files_processing_status(
                files_with_chunks,
                DocumentProcessingStatus.FAILED,
                error=f"Ошибка индексации файла: {str(exc)[:450]}",
            )
            raise
        self._session.flush()
        self._set_files_processing_status(files_with_chunks, DocumentProcessingStatus.PROCESSED)
        self._append_event_and_commit(
            run=run,
            stage=ResearchRunStage.RETRIEVAL,
            progress_percent=32,
            message="Индексация фрагментов завершена.",
            metadata={"file_count": len(files_with_chunks), "chunk_count": len(chunks)},
        )
        query_vector = retrieval_service.embed_query(research_brief)
        if len(query_vector) == 3072 and any(chunk.embedding_vector is not None for chunk in chunks):
            retrieval_results = [
                RetrievalResult(chunk=chunk, score=score)
                for chunk, score in self._repository.search_chunks_by_vector(
                    self._session,
                    user_id=run.user_id,
                    chat_session_id=run.chat_session_id,
                    query_vector=query_vector,
                    limit=retrieval_limit,
                )
            ]
        else:
            retrieval_results = retrieval_service.rank_with_query_vector(
                query_vector=query_vector,
                chunks=chunks,
                limit=retrieval_limit,
            )
        source_chunks = [result.chunk for result in retrieval_results]
        drafts = orchestrator.extract_evidence(
            research_brief=research_brief,
            fragments=retrieval_service.to_source_fragments(retrieval_results),
            limit=12,
        )
        return drafts, source_chunks

    def _set_file_processing_status(
        self,
        session_file: SessionFile,
        status: DocumentProcessingStatus,
        *,
        error: str | None = None,
        text_extracted_at: datetime | None = None,
    ) -> None:
        session_file.processing_status = status
        session_file.processing_error = error
        if text_extracted_at is not None:
            session_file.text_extracted_at = text_extracted_at
        self._session.commit()

    def _set_files_processing_status(
        self,
        files: list[SessionFile],
        status: DocumentProcessingStatus,
        *,
        error: str | None = None,
    ) -> None:
        if not files:
            return
        for session_file in files:
            session_file.processing_status = status
            session_file.processing_error = error
        self._session.commit()

    @staticmethod
    def _files_for_chunks(*, files: list[SessionFile], chunks: list[DocumentChunk]) -> list[SessionFile]:
        file_ids_with_chunks = {chunk.session_file_id for chunk in chunks}
        return [session_file for session_file in files if session_file.id in file_ids_with_chunks]

    def _persist_evidence(
        self,
        *,
        run: ResearchRun,
        evidence_drafts: list[EvidenceDraft],
        source_chunks: list[DocumentChunk],
    ) -> list[EvidenceItem]:
        evidence_items: list[EvidenceItem] = []
        for draft in evidence_drafts:
            chunk = self._chunk_by_source_index(chunks=source_chunks, source_index=draft.source_index)
            metadata = dict(draft.metadata)
            if draft.source_index is not None:
                metadata["source_index"] = draft.source_index
            relevance_score = self._evidence_relevance_score(draft=draft, chunk=chunk)
            if chunk is not None:
                metadata["source_snapshot"] = {
                    "session_file_id": str(chunk.session_file_id),
                    "chunk_id": str(chunk.id),
                    "chunk_position": chunk.position,
                    "content_hash": chunk.content_hash,
                    "source_metadata": chunk.source_metadata,
                }
            evidence_items.append(
                self._repository.create_evidence(
                    self._session,
                    EvidenceItem(
                        run_id=run.id,
                        session_file_id=chunk.session_file_id if chunk is not None else None,
                        chunk_id=chunk.id if chunk is not None else None,
                        kind=draft.kind,
                        title=draft.title,
                        summary=draft.summary,
                        quote=draft.quote,
                        confidence=draft.confidence,
                        relevance_score=relevance_score,
                        rank=len(evidence_items) + 1,
                        item_metadata=metadata,
                    ),
                )
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.EVIDENCE,
                progress_percent=44,
                message=f"Evidence сохранён: {draft.title}.",
                metadata={
                    "evidence_index": len(evidence_items) - 1,
                    "evidence_title": draft.title,
                    "chunk_id": str(chunk.id) if chunk is not None else None,
                    "session_file_id": str(chunk.session_file_id) if chunk is not None else None,
                },
            )
        return evidence_items

    def _generate_hypothesis_drafts(
        self,
        *,
        research_brief: str,
        evidence_drafts: list[EvidenceDraft],
        hypothesis_count: int,
    ) -> list[HypothesisDraft]:
        return self._require_llm_orchestrator().generate_hypotheses(
            research_brief=research_brief,
            evidence=evidence_drafts,
            hypothesis_count=hypothesis_count,
        )

    def _persist_hypotheses(
        self,
        *,
        run: ResearchRun,
        hypothesis_drafts: list[HypothesisDraft],
        evidence: list[EvidenceItem],
    ) -> list[HypothesisCandidate]:
        hypotheses: list[HypothesisCandidate] = []
        for index, draft in enumerate(hypothesis_drafts):
            hypothesis = self._repository.create_hypothesis(
                self._session,
                HypothesisCandidate(
                    run_id=run.id,
                    position=index,
                    title=draft.title,
                    statement=draft.statement,
                    mechanism=draft.mechanism,
                    kpi_alignment=draft.kpi_alignment,
                    feasibility=draft.feasibility,
                    risk_profile=draft.risk_profile,
                    novelty=draft.novelty,
                ),
            )
            hypotheses.append(hypothesis)
            for link in draft.evidence_links:
                if 0 <= link.evidence_index < len(evidence):
                    self._repository.create_hypothesis_evidence_link(
                        self._session,
                        HypothesisEvidenceLink(
                            hypothesis_id=hypothesis.id,
                            evidence_id=evidence[link.evidence_index].id,
                            relation=link.relation,
                            rationale=link.rationale,
                        ),
                    )
            self._repository.create_hypothesis_version(
                self._session,
                HypothesisVersion(
                    hypothesis_id=hypothesis.id,
                    version_number=1,
                    statement=draft.statement,
                    change_summary="Стартовая версия после генерации первичного набора гипотез.",
                    created_by_role=None,
                ),
            )
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.HYPOTHESIS_GENERATION,
                progress_percent=58,
                message=f"Гипотеза создана: {draft.title}.",
                metadata={
                    "hypothesis_id": str(hypothesis.id),
                    "hypothesis_title": draft.title,
                    "hypothesis_position": index,
                },
            )
        return hypotheses

    def _create_debate_artifacts(
        self,
        *,
        run: ResearchRun,
        hypothesis_drafts: list[HypothesisDraft],
        hypotheses: list[HypothesisCandidate],
        evidence_drafts: list[EvidenceDraft],
        round_limit: int,
    ) -> list[str]:
        refined_statements: list[str] = []
        orchestrator = self._require_llm_orchestrator()
        for hypothesis, draft in zip(hypotheses, hypothesis_drafts, strict=True):
            current_draft = draft
            current_statement = draft.statement
            self._append_event_and_commit(
                run=run,
                stage=ResearchRunStage.DEBATE,
                progress_percent=70,
                message=f"Началось обсуждение гипотезы: {hypothesis.title}.",
                metadata={"hypothesis_id": str(hypothesis.id), "hypothesis_title": hypothesis.title},
            )
            for round_number in range(1, round_limit + 1):
                debate = orchestrator.debate_hypothesis(
                    hypothesis=current_draft,
                    evidence=evidence_drafts,
                    round_number=round_number,
                )
                for role, content in debate.messages:
                    self._repository.create_debate_message(
                        self._session,
                        DebateMessage(
                            hypothesis_id=hypothesis.id,
                            round_number=round_number,
                            role=role,
                            content=content,
                        ),
                    )
                    self._append_event_and_commit(
                        run=run,
                        stage=ResearchRunStage.DEBATE,
                        progress_percent=72,
                        message=f"{self._debate_role_label(role)} ответил по гипотезе: {hypothesis.title}.",
                        metadata={
                            "hypothesis_id": str(hypothesis.id),
                            "hypothesis_title": hypothesis.title,
                            "round_number": round_number,
                            "actor_role": role.value,
                        },
                    )
                self._repository.create_hypothesis_version(
                    self._session,
                    HypothesisVersion(
                        hypothesis_id=hypothesis.id,
                        version_number=round_number + 1,
                        statement=debate.refined_statement,
                        change_summary=debate.change_summary,
                        created_by_role=DebateRole.MANUFACTURER,
                    ),
                )
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.DEBATE,
                    progress_percent=76,
                    message=f"Гипотеза доработана после {round_number} цикла: {hypothesis.title}.",
                    metadata={
                        "hypothesis_id": str(hypothesis.id),
                        "hypothesis_title": hypothesis.title,
                        "round_number": round_number,
                    },
                )
                current_statement = debate.refined_statement
                current_draft = HypothesisDraft(
                    title=current_draft.title,
                    statement=current_statement,
                    mechanism=current_draft.mechanism,
                    kpi_alignment=current_draft.kpi_alignment,
                    feasibility=current_draft.feasibility,
                    risk_profile=current_draft.risk_profile,
                    novelty=current_draft.novelty,
                    evidence_links=current_draft.evidence_links,
                )
            refined_statements.append(current_statement)
        return refined_statements

    def _create_evaluation_artifacts(
        self,
        *,
        run: ResearchRun,
        chat_session_id: uuid.UUID,
        user: User,
        hypothesis_drafts: list[HypothesisDraft],
        hypotheses: list[HypothesisCandidate],
        refined_statements: list[str],
        evidence_drafts: list[EvidenceDraft],
    ) -> list[dict[str, object]]:
        selected_agents = self._repository.list_selected_agents(
            self._session,
            user_id=user.id,
            chat_session_id=chat_session_id,
        )
        evaluators = self._evaluation_agents(selected_agents)
        orchestrator = self._require_llm_orchestrator()
        evaluation_payloads: list[dict[str, object]] = []
        for index, hypothesis in enumerate(hypotheses):
            for evaluator_key, evaluator_name, user_agent_id, evaluator_prompt in evaluators:
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.EVALUATION,
                    progress_percent=84,
                    message=f"{evaluator_name} оценивает гипотезу: {hypothesis.title}.",
                    metadata={
                        "hypothesis_id": str(hypothesis.id),
                        "hypothesis_title": hypothesis.title,
                        "evaluator_key": evaluator_key,
                        "evaluator_name": evaluator_name,
                        "user_agent_id": str(user_agent_id) if user_agent_id else None,
                    },
                )
                draft = orchestrator.evaluate_hypothesis(
                    evaluator_key=evaluator_key,
                    evaluator_name=evaluator_name,
                    evaluator_prompt=evaluator_prompt,
                    hypothesis=hypothesis_drafts[index],
                    refined_statement=refined_statements[index],
                    evidence=evidence_drafts,
                )
                self._repository.create_evaluation_result(
                    self._session,
                    EvaluationResult(
                        run_id=run.id,
                        hypothesis_id=hypothesis.id,
                        user_agent_id=user_agent_id,
                        evaluator_key=evaluator_key,
                        evaluator_name=evaluator_name,
                        score=draft.score,
                        verdict=draft.verdict,
                        rationale=draft.rationale,
                        risk_notes=draft.risk_notes,
                    ),
                )
                self._append_event_and_commit(
                    run=run,
                    stage=ResearchRunStage.EVALUATION,
                    progress_percent=88,
                    message=f"{evaluator_name} сохранил оценку гипотезы: {hypothesis.title}.",
                    metadata={
                        "hypothesis_id": str(hypothesis.id),
                        "hypothesis_title": hypothesis.title,
                        "evaluator_key": evaluator_key,
                        "evaluator_name": evaluator_name,
                        "user_agent_id": str(user_agent_id) if user_agent_id else None,
                    },
                )
                evaluation_payloads.append(
                    {
                        "hypothesis_index": index,
                        "evaluator_key": evaluator_key,
                        "evaluator_name": evaluator_name,
                        "score": float(draft.score),
                        "verdict": draft.verdict,
                        "rationale": draft.rationale,
                        "risk_notes": draft.risk_notes,
                    }
                )
        return evaluation_payloads

    def _create_judge_verdict(
        self,
        *,
        run: ResearchRun,
        research_brief: str,
        hypothesis_drafts: list[HypothesisDraft],
        refined_statements: list[str],
        evidence_drafts: list[EvidenceDraft],
        evaluation_payloads: list[dict[str, object]],
    ) -> JudgeVerdict:
        draft = self._require_llm_orchestrator().judge(
            research_brief=research_brief,
            hypotheses=hypothesis_drafts,
            refined_statements=refined_statements,
            evidence=evidence_drafts,
            evaluations=evaluation_payloads,
        )
        return self._repository.create_judge_verdict(
            self._session,
            JudgeVerdict(
                run_id=run.id,
                summary=draft.summary,
                recommendation=draft.recommendation,
                ranking=draft.ranking,
                next_checks=draft.next_checks,
            ),
        )

    def _build_detail_response(self, run: ResearchRun) -> ResearchRunDetailResponse:
        inputs = self._repository.list_input_items(self._session, run_id=run.id)
        evidence = self._repository.list_evidence(self._session, run_id=run.id)
        hypotheses = self._repository.list_hypotheses(self._session, run_id=run.id)
        hypothesis_ids = [hypothesis.id for hypothesis in hypotheses]
        versions_by_hypothesis = self._group_by(
            self._repository.list_hypothesis_versions(self._session, hypothesis_ids=hypothesis_ids),
            "hypothesis_id",
        )
        links_by_hypothesis = self._group_by(
            self._repository.list_evidence_links(self._session, hypothesis_ids=hypothesis_ids),
            "hypothesis_id",
        )
        messages_by_hypothesis = self._group_by(
            self._repository.list_debate_messages(self._session, hypothesis_ids=hypothesis_ids),
            "hypothesis_id",
        )
        evaluations_by_hypothesis = self._group_by(
            self._repository.list_evaluation_results(self._session, run_id=run.id),
            "hypothesis_id",
        )
        verdict = self._repository.get_judge_verdict(self._session, run_id=run.id)
        base = ResearchRunSummaryResponse.model_validate(run).model_dump()
        return ResearchRunDetailResponse(
            **base,
            inputs=[ResearchInputResponse.model_validate(item) for item in inputs],
            evidence=[EvidenceItemResponse.model_validate(item) for item in evidence],
            hypotheses=[
                HypothesisCandidateResponse(
                    **HypothesisCandidateResponse.model_validate(hypothesis).model_dump(exclude={
                        "versions",
                        "evidence_links",
                        "debate_messages",
                        "evaluations",
                    }),
                    versions=[
                        HypothesisVersionResponse.model_validate(item)
                        for item in versions_by_hypothesis[hypothesis.id]
                    ],
                    evidence_links=[
                        HypothesisEvidenceLinkResponse.model_validate(item)
                        for item in links_by_hypothesis[hypothesis.id]
                    ],
                    debate_messages=[
                        DebateMessageResponse.model_validate(item)
                        for item in messages_by_hypothesis[hypothesis.id]
                    ],
                    evaluations=[
                        EvaluationResultResponse.model_validate(item)
                        for item in evaluations_by_hypothesis[hypothesis.id]
                    ],
                )
                for hypothesis in hypotheses
            ],
            verdict=JudgeVerdictResponse.model_validate(verdict) if verdict is not None else None,
        )

    def _get_chat_session(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        for_update: bool = False,
    ):
        chat_session = self._repository.get_chat_session(
            self._session,
            user_id=user.id,
            chat_session_id=chat_session_id,
            for_update=for_update,
        )
        if chat_session is None:
            raise NotFoundError("Чат не найден.")
        return chat_session

    def _get_run(self, *, user: User, chat_session_id: uuid.UUID, run_id: uuid.UUID) -> ResearchRun:
        self._get_chat_session(user=user, chat_session_id=chat_session_id)
        run = self._repository.get_run(
            self._session,
            user_id=user.id,
            chat_session_id=chat_session_id,
            run_id=run_id,
        )
        if run is None:
            raise NotFoundError("Версия исследования не найдена.")
        return run

    def _mark_run_failed(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        error: Exception,
    ) -> None:
        try:
            run = self._repository.get_run(
                self._session,
                user_id=user.id,
                chat_session_id=chat_session_id,
                run_id=run_id,
                for_update=True,
            )
            if run is None or run.status != ResearchRunStatus.RUNNING:
                self._session.rollback()
                return
            run.status = ResearchRunStatus.FAILED
            run.completed_at = datetime.now(UTC)
            run.error_message = self._pipeline_error_message(error)
            self._append_event(
                run=run,
                stage=ResearchRunStage.FAILED,
                progress_percent=100,
                message=run.error_message,
            )
            self._session.commit()
        except Exception:
            self._session.rollback()

    def _mark_run_cancelled(
        self,
        *,
        user: User,
        chat_session_id: uuid.UUID,
        run_id: uuid.UUID,
        message: str,
    ) -> None:
        try:
            run = self._repository.get_run(
                self._session,
                user_id=user.id,
                chat_session_id=chat_session_id,
                run_id=run_id,
                for_update=True,
            )
            if run is None:
                self._session.rollback()
                return
            if run.status == ResearchRunStatus.CANCELLED:
                self._session.rollback()
                return
            run.status = ResearchRunStatus.CANCELLED
            run.completed_at = datetime.now(UTC)
            self._append_event(
                run=run,
                stage=ResearchRunStage.CANCELLED,
                progress_percent=100,
                message=message,
            )
            self._session.commit()
        except Exception:
            self._session.rollback()

    @staticmethod
    def _normalize_inputs(input_requests: list[ResearchInputRequest]) -> list[NormalizedInput]:
        normalized: list[NormalizedInput] = []
        for position, item in enumerate(input_requests):
            label = item.label.strip()
            text = item.text.strip()
            if not label:
                raise ValidationError("Название поля ввода не может быть пустым.")
            if len(label) > 64:
                raise ValidationError("Название поля ввода должно быть не длиннее 64 символов.")
            if not text:
                raise ValidationError("Текст поля ввода не может быть пустым.")
            if len(text) > 4000:
                raise ValidationError("Текст одного поля ввода должен быть не длиннее 4000 символов.")
            normalized.append(NormalizedInput(kind=item.kind, label=label, text=text, position=position))
        return normalized

    @staticmethod
    def _create_input_hash(inputs: list[NormalizedInput]) -> str:
        payload = [
            {"kind": item.kind.value, "label": item.label, "text": item.text, "position": item.position}
            for item in inputs
        ]
        raw = json.dumps(payload, ensure_ascii=False, sort_keys=True).encode("utf-8")
        return hashlib.sha256(raw).hexdigest()

    @staticmethod
    def _create_title(inputs: list[NormalizedInput]) -> str:
        source = next((item.text for item in inputs if item.kind == ResearchInputKind.CONTEXT), inputs[0].text)
        words = re.sub(r"[.,;:!?]+$", "", source).split()
        title = " ".join(words[:7]).strip()
        if not title:
            return "Новая проверка гипотез"
        return title[:197].rstrip() + "..." if len(title) > 200 else title

    @staticmethod
    def _short_text(text: str, limit: int) -> str:
        normalized = _WORD_RE.sub(" ", text).strip()
        if len(normalized) <= limit:
            return normalized
        return normalized[: max(0, limit - 3)].rstrip() + "..."

    def _append_event(
        self,
        *,
        run: ResearchRun,
        stage: ResearchRunStage,
        progress_percent: int,
        message: str,
        metadata: dict[str, object] | None = None,
    ) -> ResearchRunEvent:
        return self._repository.create_run_event(
            self._session,
            ResearchRunEvent(
                run_id=run.id,
                stage=stage,
                sequence_number=self._repository.next_event_sequence_number(self._session, run_id=run.id),
                progress_percent=progress_percent,
                message=self._short_text(message, 500),
                event_metadata=metadata or {},
                created_at=datetime.now(UTC),
            ),
        )

    def _append_event_and_commit(
        self,
        *,
        run: ResearchRun,
        stage: ResearchRunStage,
        progress_percent: int,
        message: str,
        metadata: dict[str, object] | None = None,
    ) -> ResearchRunEvent:
        event = self._append_event(
            run=run,
            stage=stage,
            progress_percent=progress_percent,
            message=message,
            metadata=metadata,
        )
        self._session.commit()
        return event

    def _check_cancelled(self, run: ResearchRun) -> None:
        if hasattr(self._session, "expire"):
            self._session.expire(run, ["status"])
            current_status = run.status
        else:
            refreshed_run = self._repository.get_run(
                self._session,
                user_id=run.user_id,
                chat_session_id=run.chat_session_id,
                run_id=run.id,
            )
            current_status = refreshed_run.status if refreshed_run is not None else run.status
        if current_status == ResearchRunStatus.CANCELLED:
            raise ConflictError("Исследовательский запуск отменён.")

    @staticmethod
    def _pipeline_error_message(error: Exception) -> str:
        if isinstance(error, ServiceError):
            return ResearchService._short_text(error.detail, 2000)
        return "Внутренняя ошибка исследовательского пайплайна"

    @staticmethod
    def _debate_role_label(role: DebateRole) -> str:
        return {
            DebateRole.DEFENDER: "Защитник",
            DebateRole.ATTACKER: "Атакующий",
            DebateRole.MANUFACTURER: "Производственник",
        }.get(role, role.value)

    @staticmethod
    def _apply_pipeline_settings_to_brief(research_brief: str, pipeline_settings) -> str:
        additions: list[str] = []
        if pipeline_settings.excluded_directions:
            additions.append(f"Исключённые направления: {pipeline_settings.excluded_directions}")
        if pipeline_settings.domain_constraints:
            additions.append(f"Доменные ограничения: {pipeline_settings.domain_constraints}")
        if not additions:
            return research_brief
        return research_brief + "\n" + "\n".join(additions)

    @staticmethod
    def _evidence_relevance_score(*, draft: EvidenceDraft, chunk: DocumentChunk | None) -> Decimal | None:
        if chunk is None:
            return None
        raw_score = draft.metadata.get("retrieval_score") if draft.metadata else None
        try:
            score = Decimal(str(raw_score)) if raw_score is not None else draft.confidence
        except Exception:
            score = draft.confidence
        return max(Decimal("0.000000"), min(Decimal("1.000000"), score)).quantize(Decimal("0.000001"))

    @staticmethod
    def _research_brief(inputs: list[NormalizedInput]) -> str:
        return "\n".join(f"{item.label} ({item.kind.value}): {item.text}" for item in inputs)

    @staticmethod
    def _chunk_by_source_index(*, chunks: list[DocumentChunk], source_index: int | None) -> DocumentChunk | None:
        if source_index is None or source_index < 0 or source_index >= len(chunks):
            return None
        return chunks[source_index]

    def _validate_feedback_target(self, *, run: ResearchRun, payload: ResearchFeedbackCreateRequest) -> None:
        if payload.target_type == ResearchFeedbackTarget.RUN:
            if payload.target_id is not None and payload.target_id != run.id:
                raise ValidationError("Feedback на запуск должен ссылаться на текущую версию исследования.")
            return
        if payload.target_id is None:
            raise ValidationError("Для feedback по гипотезе или вердикту нужен target_id.")
        if payload.target_type == ResearchFeedbackTarget.HYPOTHESIS:
            hypothesis_ids = {item.id for item in self._repository.list_hypotheses(self._session, run_id=run.id)}
            if payload.target_id not in hypothesis_ids:
                raise NotFoundError("Гипотеза для feedback не найдена.")
            return
        if payload.target_type == ResearchFeedbackTarget.VERDICT:
            verdict = self._repository.get_judge_verdict(self._session, run_id=run.id)
            if verdict is None or payload.target_id != verdict.id:
                raise NotFoundError("Вердикт для feedback не найден.")

    @staticmethod
    def _normalize_optional_text(value: str | None) -> str | None:
        if value is None:
            return None
        normalized = _WORD_RE.sub(" ", value).strip()
        return normalized or None

    @staticmethod
    def _build_markdown_export(detail: ResearchRunDetailResponse) -> str:
        lines = [
            f"# {detail.title}",
            "",
            f"- Версия: {detail.version_number}",
            f"- Статус: {detail.status.value}",
            f"- Модель: {detail.model_name or 'не указана'}",
            "",
            "## Входные параметры",
        ]
        for item in detail.inputs:
            lines.append(f"- **{item.label}** ({item.kind.value}): {item.text}")
        lines.extend(["", "## Evidence"])
        for item in detail.evidence:
            lines.append(f"- **{item.title}** [{item.kind.value}]: {item.summary}")
        lines.extend(["", "## Гипотезы"])
        for hypothesis in detail.hypotheses:
            lines.append(f"### {hypothesis.title}")
            lines.append(hypothesis.statement)
            lines.append(f"- Механизм: {hypothesis.mechanism}")
            lines.append(f"- Риски: {hypothesis.risk_profile}")
        if detail.verdict is not None:
            lines.extend(["", "## Вердикт", detail.verdict.summary, "", detail.verdict.recommendation])
            if detail.verdict.next_checks:
                lines.extend(["", "## Первые проверки"])
                lines.extend(f"- {item}" for item in detail.verdict.next_checks)
        return "\n".join(lines).strip() + "\n"

    @staticmethod
    def _build_csv_export(detail: ResearchRunDetailResponse) -> str:
        output = io.StringIO()
        writer = csv.writer(output)
        writer.writerow(["section", "id", "parent_id", "field", "value"])
        writer.writerow(["run", str(detail.id), "", "title", ResearchService._csv_safe(detail.title)])
        writer.writerow(["run", str(detail.id), "", "version", detail.version_number])
        writer.writerow(["run", str(detail.id), "", "status", detail.status.value])
        for item in detail.inputs:
            writer.writerow(["input", str(item.id), str(detail.id), ResearchService._csv_safe(item.label), ResearchService._csv_safe(item.text)])
        for item in detail.evidence:
            writer.writerow(["evidence", str(item.id), str(detail.id), item.kind.value, ResearchService._csv_safe(item.summary)])
            if item.quote:
                writer.writerow(["evidence", str(item.id), str(detail.id), "quote", ResearchService._csv_safe(item.quote)])
        for hypothesis in detail.hypotheses:
            writer.writerow(
                [
                    "hypothesis",
                    str(hypothesis.id),
                    str(detail.id),
                    ResearchService._csv_safe(hypothesis.title),
                    ResearchService._csv_safe(hypothesis.statement),
                ]
            )
            writer.writerow(["hypothesis", str(hypothesis.id), str(detail.id), "mechanism", ResearchService._csv_safe(hypothesis.mechanism)])
            writer.writerow(["hypothesis", str(hypothesis.id), str(detail.id), "risk_profile", ResearchService._csv_safe(hypothesis.risk_profile)])
            for evaluation in hypothesis.evaluations:
                writer.writerow(
                    [
                        "evaluation",
                        str(evaluation.id),
                        str(hypothesis.id),
                        ResearchService._csv_safe(evaluation.evaluator_name),
                        ResearchService._csv_safe(f"{evaluation.score}: {evaluation.verdict}"),
                    ]
                )
        if detail.verdict is not None:
            writer.writerow(["verdict", str(detail.verdict.id), str(detail.id), "summary", ResearchService._csv_safe(detail.verdict.summary)])
            writer.writerow(
                [
                    "verdict",
                    str(detail.verdict.id),
                    str(detail.id),
                    "recommendation",
                    ResearchService._csv_safe(detail.verdict.recommendation),
                ]
            )
            for index, check in enumerate(detail.verdict.next_checks, start=1):
                writer.writerow(
                    [
                        "next_check",
                        f"{detail.verdict.id}:{index}",
                        str(detail.verdict.id),
                        str(index),
                        ResearchService._csv_safe(check),
                    ]
                )
        return output.getvalue()

    @staticmethod
    def _csv_safe(value: object) -> object:
        if not isinstance(value, str) or not value:
            return value
        if value[0] in ("=", "+", "-", "@", "\t", "\r", "\n"):
            return "'" + value
        return value

    @staticmethod
    def _build_docx_export(detail: ResearchRunDetailResponse) -> str:
        from docx import Document

        document = Document()
        document.add_heading(detail.title, level=1)
        document.add_paragraph(f"Версия: {detail.version_number}")
        document.add_paragraph(f"Статус: {detail.status.value}")
        document.add_paragraph(f"Модель: {detail.model_name or 'не указана'}")

        document.add_heading("Входные параметры", level=2)
        for item in detail.inputs:
            document.add_paragraph(f"{item.label} ({item.kind.value}): {item.text}", style="List Bullet")

        document.add_heading("Evidence", level=2)
        for item in detail.evidence:
            document.add_paragraph(f"{item.title} [{item.kind.value}]", style="List Bullet")
            document.add_paragraph(item.summary)
            if item.quote:
                document.add_paragraph(f"Цитата: {item.quote}")

        document.add_heading("Гипотезы", level=2)
        for hypothesis in detail.hypotheses:
            document.add_heading(hypothesis.title, level=3)
            document.add_paragraph(hypothesis.statement)
            document.add_paragraph(f"Механизм: {hypothesis.mechanism}")
            document.add_paragraph(f"Риски: {hypothesis.risk_profile}")
            if hypothesis.evaluations:
                document.add_paragraph("Оценки:")
                for evaluation in hypothesis.evaluations:
                    document.add_paragraph(
                        f"{evaluation.evaluator_name}: {evaluation.score} - {evaluation.verdict}",
                        style="List Bullet",
                    )

        if detail.verdict is not None:
            document.add_heading("Вердикт", level=2)
            document.add_paragraph(detail.verdict.summary)
            document.add_paragraph(detail.verdict.recommendation)
            if detail.verdict.next_checks:
                document.add_heading("Первые проверки", level=2)
                for item in detail.verdict.next_checks:
                    document.add_paragraph(item, style="List Bullet")

        buffer = io.BytesIO()
        document.save(buffer)
        return base64.b64encode(buffer.getvalue()).decode("ascii")

    @staticmethod
    def _build_pdf_export(detail: ResearchRunDetailResponse) -> str:
        from reportlab.lib.pagesizes import A4
        from reportlab.lib.styles import getSampleStyleSheet
        from reportlab.platypus import Paragraph, SimpleDocTemplate, Spacer
        from reportlab.pdfbase import pdfmetrics
        from reportlab.pdfbase.ttfonts import TTFont

        font_name = "Helvetica"
        for font_path in (
            "/usr/share/fonts/truetype/dejavu/DejaVuSans.ttf",
            "/usr/local/share/fonts/DejaVuSans.ttf",
        ):
            try:
                pdfmetrics.registerFont(TTFont("DejaVuSans", font_path))
                font_name = "DejaVuSans"
                break
            except Exception:
                continue

        buffer = io.BytesIO()
        document = SimpleDocTemplate(buffer, pagesize=A4, title=detail.title)
        styles = getSampleStyleSheet()
        for style_name in ("Title", "Heading1", "Heading2", "Heading3", "BodyText", "Bullet"):
            styles[style_name].fontName = font_name

        story = [
            Paragraph(escape(detail.title), styles["Title"]),
            Paragraph(f"Версия: {detail.version_number}", styles["BodyText"]),
            Paragraph(f"Статус: {escape(detail.status.value)}", styles["BodyText"]),
            Paragraph(f"Модель: {escape(detail.model_name or 'не указана')}", styles["BodyText"]),
            Spacer(1, 12),
            Paragraph("Входные параметры", styles["Heading2"]),
        ]
        for item in detail.inputs:
            story.append(Paragraph(escape(f"{item.label} ({item.kind.value}): {item.text}"), styles["Bullet"]))
        story.extend([Spacer(1, 12), Paragraph("Evidence", styles["Heading2"])])
        for item in detail.evidence:
            story.append(Paragraph(escape(f"{item.title} [{item.kind.value}]: {item.summary}"), styles["BodyText"]))
        story.extend([Spacer(1, 12), Paragraph("Гипотезы", styles["Heading2"])])
        for hypothesis in detail.hypotheses:
            story.append(Paragraph(escape(hypothesis.title), styles["Heading3"]))
            story.append(Paragraph(escape(hypothesis.statement), styles["BodyText"]))
            story.append(Paragraph(escape(f"Риски: {hypothesis.risk_profile}"), styles["BodyText"]))
        if detail.verdict is not None:
            story.extend([Spacer(1, 12), Paragraph("Вердикт", styles["Heading2"])])
            story.append(Paragraph(escape(detail.verdict.summary), styles["BodyText"]))
            story.append(Paragraph(escape(detail.verdict.recommendation), styles["BodyText"]))
            if detail.verdict.next_checks:
                story.append(Paragraph("Первые проверки", styles["Heading2"]))
                for item in detail.verdict.next_checks:
                    story.append(Paragraph(escape(item), styles["Bullet"]))

        document.build(story)
        return base64.b64encode(buffer.getvalue()).decode("ascii")

    def _require_llm_orchestrator(self) -> ResearchLlmOrchestrator:
        if self._llm_orchestrator is None:
            raise ValidationError("LLM-пайплайн исследования не настроен")
        return self._llm_orchestrator

    def _require_retrieval_service(self) -> ResearchRetrievalService:
        if self._retrieval_service is None:
            raise ValidationError("Retrieval-пайплайн исследования не настроен")
        return self._retrieval_service

    @staticmethod
    def _evaluation_agents(selected_agents: list[UserAgent]) -> list[tuple[str, str, uuid.UUID | None, str | None]]:
        if selected_agents:
            return [(str(agent.id), agent.name, agent.id, agent.system_prompt) for agent in selected_agents]
        return [
            ("finance", "Финансовый агент", None, "Оцени экономическую целесообразность, стоимость проверки и потенциал эффекта."),
            ("risk", "Риск-агент", None, "Оцени неопределённость, вероятность неудачи и критичные условия провала."),
        ]

    @staticmethod
    def _group_by(items, attribute: str):
        grouped = defaultdict(list)
        for item in items:
            grouped[getattr(item, attribute)].append(item)
        return grouped
