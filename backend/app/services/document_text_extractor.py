import csv
import io
import json
import zipfile
from dataclasses import dataclass
from pathlib import Path
from xml.etree import ElementTree

from bs4 import BeautifulSoup
from docx import Document
from openpyxl import load_workbook
from pypdf import PdfReader
from pypdf.errors import PdfReadError

from app.services.exceptions import ValidationError

_TEXT_FILE_KINDS = {"txt", "csv", "json", "md", "markdown", "xml", "html", "htm", "log"}
_MAX_EXTRACT_BYTES = 15 * 1024 * 1024
_XML_TEXT_LIMIT = 2 * 1024 * 1024


@dataclass(frozen=True)
class ExtractedText:
    text: str
    metadata: dict[str, object]


class DocumentTextExtractor:
    def extract(self, *, path: Path, kind: str | None, content_type: str | None = None) -> ExtractedText | None:
        if not path.exists():
            raise ValidationError("Файл источника не найден")
        if path.stat().st_size > _MAX_EXTRACT_BYTES:
            raise ValidationError("Файл слишком большой для синхронной обработки")

        normalized_kind = self._normalize_kind(path=path, kind=kind, content_type=content_type)
        if normalized_kind in {"txt", "md", "markdown", "log"}:
            return ExtractedText(text=self._read_plain_text(path), metadata={"parser": "plain_text"})
        if normalized_kind == "csv":
            return ExtractedText(text=self._read_csv(path), metadata={"parser": "csv"})
        if normalized_kind == "json":
            return ExtractedText(text=self._read_json(path), metadata={"parser": "json"})
        if normalized_kind == "xml":
            return ExtractedText(text=self._read_xml(path), metadata={"parser": "xml"})
        if normalized_kind in {"html", "htm"}:
            return ExtractedText(text=self._read_html(path), metadata={"parser": "html"})
        if normalized_kind == "pdf":
            return self._read_pdf(path)
        if normalized_kind == "docx":
            return self._read_docx(path)
        if normalized_kind == "xlsx":
            return self._read_xlsx(path)
        return None

    @staticmethod
    def supported_kinds() -> set[str]:
        return _TEXT_FILE_KINDS | {"pdf", "docx", "xlsx"}

    @staticmethod
    def _normalize_kind(*, path: Path, kind: str | None, content_type: str | None) -> str:
        if kind:
            return kind.lower().lstrip(".")
        suffix = path.suffix.lower().lstrip(".")
        if suffix:
            return suffix
        if content_type == "application/pdf":
            return "pdf"
        return ""

    @staticmethod
    def _read_plain_text(path: Path) -> str:
        return path.read_text(encoding="utf-8").strip()

    @staticmethod
    def _read_csv(path: Path) -> str:
        raw = path.read_text(encoding="utf-8-sig")
        rows = csv.reader(io.StringIO(raw))
        lines = [" | ".join(cell.strip() for cell in row if cell.strip()) for row in rows]
        return "\n".join(line for line in lines if line).strip()

    @staticmethod
    def _read_json(path: Path) -> str:
        try:
            payload = json.loads(path.read_text(encoding="utf-8"))
        except json.JSONDecodeError as exc:
            raise ValidationError("JSON-файл повреждён или имеет неподдерживаемый формат") from exc
        return json.dumps(payload, ensure_ascii=False, indent=2)

    @staticmethod
    def _read_xml(path: Path) -> str:
        raw = path.read_text(encoding="utf-8")
        if len(raw) > _XML_TEXT_LIMIT:
            raise ValidationError("XML-файл слишком большой для синхронной обработки")
        try:
            root = ElementTree.fromstring(raw)
        except ElementTree.ParseError as exc:
            raise ValidationError("XML-файл повреждён или имеет неподдерживаемый формат") from exc
        fragments = [text.strip() for text in root.itertext() if text and text.strip()]
        return "\n".join(fragments)

    @staticmethod
    def _read_html(path: Path) -> str:
        soup = BeautifulSoup(path.read_text(encoding="utf-8"), "html.parser")
        for tag in soup(["script", "style", "noscript"]):
            tag.extract()
        return soup.get_text(separator="\n", strip=True)

    @staticmethod
    def _read_pdf(path: Path) -> ExtractedText:
        try:
            reader = PdfReader(str(path))
        except PdfReadError as exc:
            raise ValidationError("PDF-файл повреждён или имеет неподдерживаемый формат") from exc
        page_texts: list[str] = []
        for index, page in enumerate(reader.pages, start=1):
            text = (page.extract_text() or "").strip()
            if text:
                page_texts.append(f"[Страница {index}]\n{text}")
        return ExtractedText(
            text="\n\n".join(page_texts).strip(),
            metadata={"parser": "pdf", "pages": len(reader.pages)},
        )

    @staticmethod
    def _read_docx(path: Path) -> ExtractedText:
        try:
            document = Document(str(path))
        except (ValueError, zipfile.BadZipFile) as exc:
            raise ValidationError("DOCX-файл повреждён или имеет неподдерживаемый формат") from exc
        paragraphs = [paragraph.text.strip() for paragraph in document.paragraphs if paragraph.text.strip()]
        table_rows: list[str] = []
        for table in document.tables:
            for row in table.rows:
                cells = [cell.text.strip() for cell in row.cells if cell.text.strip()]
                if cells:
                    table_rows.append(" | ".join(cells))
        parts = paragraphs + table_rows
        return ExtractedText(
            text="\n".join(parts).strip(),
            metadata={"parser": "docx", "paragraphs": len(paragraphs), "table_rows": len(table_rows)},
        )

    @staticmethod
    def _read_xlsx(path: Path) -> ExtractedText:
        try:
            workbook = load_workbook(filename=str(path), read_only=True, data_only=True)
        except (ValueError, zipfile.BadZipFile) as exc:
            raise ValidationError("XLSX-файл повреждён или имеет неподдерживаемый формат") from exc
        lines: list[str] = []
        sheet_count = 0
        for sheet in workbook.worksheets:
            sheet_count += 1
            lines.append(f"[Лист: {sheet.title}]")
            for row in sheet.iter_rows(values_only=True):
                cells = [str(value).strip() for value in row if value is not None and str(value).strip()]
                if cells:
                    lines.append(" | ".join(cells))
        workbook.close()
        return ExtractedText(text="\n".join(lines).strip(), metadata={"parser": "xlsx", "sheets": sheet_count})
