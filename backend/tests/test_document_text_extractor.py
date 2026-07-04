from pathlib import Path

import pytest
from docx import Document
from openpyxl import Workbook

from app.services.document_text_extractor import DocumentTextExtractor
from app.services.exceptions import ValidationError


def test_extract_json_normalizes_structured_content(tmp_path: Path):
    path = tmp_path / "source.json"
    path.write_text('{"goal": "повысить прочность", "risk": "масштабирование"}', encoding="utf-8")

    extracted = DocumentTextExtractor().extract(path=path, kind="json")

    assert extracted is not None
    assert "повысить прочность" in extracted.text
    assert extracted.metadata["parser"] == "json"


def test_extract_docx_reads_paragraphs_and_tables(tmp_path: Path):
    path = tmp_path / "report.docx"
    document = Document()
    document.add_paragraph("Основной вывод по прочности.")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "Параметр"
    table.rows[0].cells[1].text = "Значение"
    document.save(path)

    extracted = DocumentTextExtractor().extract(path=path, kind="docx")

    assert extracted is not None
    assert "Основной вывод по прочности." in extracted.text
    assert "Параметр | Значение" in extracted.text
    assert extracted.metadata["parser"] == "docx"


def test_extract_xlsx_reads_sheet_rows(tmp_path: Path):
    path = tmp_path / "experiments.xlsx"
    workbook = Workbook()
    worksheet = workbook.active
    worksheet.title = "Опыт"
    worksheet.append(["KPI", "Прочность"])
    worksheet.append(["Результат", 12])
    workbook.save(path)
    workbook.close()

    extracted = DocumentTextExtractor().extract(path=path, kind="xlsx")

    assert extracted is not None
    assert "[Лист: Опыт]" in extracted.text
    assert "KPI | Прочность" in extracted.text
    assert extracted.metadata["parser"] == "xlsx"


def test_extract_broken_json_raises_validation_error(tmp_path: Path):
    path = tmp_path / "broken.json"
    path.write_text("{broken", encoding="utf-8")

    with pytest.raises(ValidationError):
        DocumentTextExtractor().extract(path=path, kind="json")
